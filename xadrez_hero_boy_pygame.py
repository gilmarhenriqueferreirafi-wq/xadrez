"""
XADREZ: HERO x BOY  (Machine Learning + tabuleiro grafico com PYGAME)
=====================================================================
Duas IAs jogam xadrez uma contra a outra numa janela do pygame e aprendem sozinhas
com o resultado de cada partida (computacao cognitiva: reconhecem padroes nos dados).

  HERO = REDE NEURAL (769 entradas -> 32 neuronios -> nota da posicao)
  BOY  = REGRESSAO com 14 padroes (material, mobilidade, centro, roque...)

Como aprendem: cada lance vira uma experiencia (posicao + resultado final da partida);
as experiencias vao para um dataset compartilhado; depois de cada partida os dois modelos
treinam em lotes sorteados desse dataset (backpropagation + Adam). A exploracao de lances
novos diminui com o tempo. Tabela compartilhada de vitorias/empates/derrotas por lance.

Tudo e salvo no Excel  memoria_xadrez_ml.xlsx  (mesma pasta; e o mesmo arquivo da versao
sem janela, entao eles continuam de onde pararam).
Hero e Boy alternam quem joga de brancas a cada partida.

CONTROLES (teclado ou clique nos botoes), valem a qualquer momento:
  A = automatico        R = rapido (sem pausas)
  L = ESPACO avanca um lance        P = ESPACO avanca uma partida
  ESPACO (no automatico/rapido) = pausa        + / - = velocidade do automatico
  G = troca o estilo das pecas      Q ou ESC = salva e sai

Instalar:  python -m pip install chess openpyxl numpy pygame-ce
           (no Python 3.14 use  pygame-ce , o  pygame  comum ainda nao tem versao para ele)
Rodar:     python xadrez_hero_boy_pygame.py          (ou:  python xadrez_hero_boy_pygame.py 50
                                                      para parar sozinho apos 50 partidas)
IMPORTANTE: feche o Excel antes de rodar, senao o arquivo nao consegue ser salvo.
"""
import os
import random
import sys
import time
from datetime import datetime

import chess
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

try:
    import pygame
except ImportError:
    print("Falta o pygame. Instale com:  python -m pip install pygame-ce")
    sys.exit(1)

PASTA = os.path.dirname(os.path.abspath(__file__))
ARQUIVO = os.path.join(PASTA, "memoria_xadrez_ml.xlsx")

MAX_LANCES = 160          # meio-lances; passou disso a partida e encerrada
PAUSA_AUTO = 0.4          # segundos entre lances no modo automatico
PESO_COMPARTILHADA = 0.3  # influencia da tabela compartilhada na escolha do lance

# --- Machine Learning ---
N_ENTRADA = 769           # 12 tipos de peca x 64 casas + material
OCULTA = 32               # neuronios do Hero
N_PADROES = 14            # padroes do Boy
GAMMA = 0.97              # desconto: lances proximos do fim valem mais
BUFFER_MAX = 3000         # tamanho do dataset (experiencias mais recentes)
LOTE = 64                 # tamanho do lote de treino
PASSOS_TREINO = 40        # lotes por partida, por modelo
MIN_TREINO = 64           # so treina quando houver experiencias suficientes
LR_HERO = 0.004
LR_BOY = 0.01

VALORES = {chess.PAWN: 1, chess.KNIGHT: 3, chess.BISHOP: 3,
           chess.ROOK: 5, chess.QUEEN: 9, chess.KING: 0}
CENTRO = (chess.D4, chess.E4, chess.D5, chess.E5)

MOTIVOS = {
    "CHECKMATE": "xeque-mate",
    "STALEMATE": "afogamento",
    "INSUFFICIENT_MATERIAL": "material insuficiente",
    "SEVENTYFIVE_MOVES": "regra dos 75 lances",
    "FIFTY_MOVES": "regra dos 50 lances",
    "FIVEFOLD_REPETITION": "repeticao (5x)",
    "THREEFOLD_REPETITION": "repeticao (3x)",
    "LIMITE_MATERIAL": "limite de lances (vantagem de material)",
    "LIMITE_EMPATE": "limite de lances (igualdade)",
}
NOMES_MODO = {
    "auto": "AUTOMATICO",
    "rapido": "RAPIDO (sem tabuleiro)",
    "lance": "MANUAL - ESPACO a cada lance",
    "partida": "MANUAL - ESPACO a cada partida",
}
TECLAS_MODO = {"a": "auto", "r": "rapido", "l": "lance", "p": "partida"}

NOMES_PADROES = [
    "diferenca de peoes", "diferenca de cavalos", "diferenca de bispos",
    "diferenca de torres", "diferenca de damas", "par de bispos",
    "mobilidade do adversario", "minha mobilidade", "controle do centro",
    "avanco dos peoes", "xeque", "rei roqueado", "desenvolvimento das pecas",
    "vies (constante)",
]
NOMES_PECAS = ["peao", "cavalo", "bispo", "torre", "dama", "rei"]


# ---------------------------------------------------------------------------
# Tabuleiro -> numeros (o que os modelos "enxergam")
# ---------------------------------------------------------------------------
def material(board, cor):
    """Saldo de material do ponto de vista de 'cor'."""
    total = 0
    for peca, valor in VALORES.items():
        total += valor * len(board.pieces(peca, cor))
        total -= valor * len(board.pieces(peca, not cor))
    return total


def feats_rede(board, cor):
    """Entrada da rede neural: 768 casas (minhas pecas / pecas do adversario) + material.
    O tabuleiro e espelhado para as pretas, entao o modelo sempre 've' do ponto de vista de quem joga."""
    x = np.zeros(N_ENTRADA, dtype=np.float32)
    for casa, p in board.piece_map().items():
        if cor == chess.BLACK:
            casa = chess.square_mirror(casa)
        grupo = (0 if p.color == cor else 6) + p.piece_type - 1
        x[grupo * 64 + casa] = 1.0
    x[768] = material(board, cor) / 10.0
    return x


def feats_padroes(board, cor):
    """Os 14 padroes que o Boy avalia (todos do ponto de vista de 'cor')."""
    meu, adv = cor, not cor
    f = np.zeros(N_PADROES, dtype=np.float32)
    escalas = (4, 2, 2, 2, 1)
    tipos = (chess.PAWN, chess.KNIGHT, chess.BISHOP, chess.ROOK, chess.QUEEN)
    for i, t in enumerate(tipos):
        f[i] = (len(board.pieces(t, meu)) - len(board.pieces(t, adv))) / escalas[i]
    f[5] = (len(board.pieces(chess.BISHOP, meu)) >= 2) - (len(board.pieces(chess.BISHOP, adv)) >= 2)

    # mobilidade (lances possiveis) de quem esta na vez e de quem nao esta
    vez = board.pseudo_legal_moves.count() / 40
    board.push(chess.Move.null())
    fora = board.pseudo_legal_moves.count() / 40
    board.pop()
    f[6], f[7] = (vez, fora) if board.turn == adv else (fora, vez)

    centro = 0
    for casa in CENTRO:
        p = board.piece_at(casa)
        if p:
            centro += 1 if p.color == meu else -1
    f[8] = centro / 4

    def avanco(c):
        total = 0
        for casa in board.pieces(chess.PAWN, c):
            r = chess.square_rank(casa)
            total += r if c == chess.WHITE else 7 - r
        return total
    f[9] = (avanco(meu) - avanco(adv)) / 16

    if board.is_check():
        f[10] = 1.0 if board.turn == adv else -1.0

    def roque(c):
        rei = board.king(c)
        if rei is None:
            return 0
        fundo = 0 if c == chess.WHITE else 7
        return int(chess.square_rank(rei) == fundo and chess.square_file(rei) in (2, 6))
    f[11] = roque(meu) - roque(adv)

    def desenvolvidas(c):
        fundo = 0 if c == chess.WHITE else 7
        n = 0
        for t in (chess.KNIGHT, chess.BISHOP):
            for casa in board.pieces(t, c):
                if chess.square_rank(casa) != fundo:
                    n += 1
        return n
    f[12] = (desenvolvidas(meu) - desenvolvidas(adv)) / 4
    f[13] = 1.0
    return f


# ---------------------------------------------------------------------------
# Modelos de Machine Learning (so numpy)
# ---------------------------------------------------------------------------
class Adam:
    """Otimizador: ajusta os pesos na direcao que reduz o erro."""

    def __init__(self, params, lr):
        self.params = params
        self.lr = lr
        self.m = [np.zeros_like(p) for p in params]
        self.v = [np.zeros_like(p) for p in params]
        self.t = 0

    def passo(self, grads):
        self.t += 1
        for p, g, m, v in zip(self.params, grads, self.m, self.v):
            m *= 0.9
            m += 0.1 * g
            v *= 0.999
            v += 0.001 * g * g
            mh = m / (1 - 0.9 ** self.t)
            vh = v / (1 - 0.999 ** self.t)
            p -= self.lr * mh / (np.sqrt(vh) + 1e-8)


class RedeNeural:
    """769 entradas -> 32 neuronios (tanh) -> 1 saida (tanh) = nota da posicao entre -1 e +1."""

    def __init__(self):
        rng = np.random.default_rng()
        self.W1 = rng.normal(0, 0.05, (N_ENTRADA, OCULTA))
        self.W1[-1, :] = np.abs(rng.normal(1.5, 0.3, OCULTA))  # ponto de partida: material importa
        self.b1 = np.zeros(OCULTA)
        self.W2 = rng.normal(0.05, 0.05, (OCULTA, 1))
        self.b2 = np.zeros(1)
        self.otimizador = Adam([self.W1, self.b1, self.W2, self.b2], LR_HERO)

    def prever(self, X):
        h = np.tanh(X @ self.W1 + self.b1)
        return np.tanh(h @ self.W2 + self.b2)[:, 0]

    def treinar(self, X, y):
        n = len(X)
        h = np.tanh(X @ self.W1 + self.b1)
        saida = np.tanh(h @ self.W2 + self.b2)[:, 0]
        erro = saida - y
        dz2 = (2 * erro / n * (1 - saida ** 2))[:, None]
        dW2 = h.T @ dz2
        db2 = dz2.sum(axis=0)
        dz1 = (dz2 @ self.W2.T) * (1 - h ** 2)
        dW1 = X.T @ dz1
        db1 = dz1.sum(axis=0)
        self.otimizador.passo([dW1, db1, dW2, db2])
        return float(np.mean(erro ** 2))


class Regressao:
    """14 padroes x pesos -> tanh -> nota da posicao entre -1 e +1."""

    def __init__(self):
        self.w = np.zeros(N_PADROES)
        # ponto de partida: valor basico das pecas (o resto ele descobre sozinho)
        self.w[:5] = [0.4, 0.6, 0.6, 1.0, 0.9]
        self.otimizador = Adam([self.w], LR_BOY)

    def prever(self, F):
        return np.tanh(F @ self.w)

    def treinar(self, F, y):
        n = len(F)
        saida = np.tanh(F @ self.w)
        erro = saida - y
        dz = 2 * erro / n * (1 - saida ** 2)
        self.otimizador.passo([F.T @ dz])
        return float(np.mean(erro ** 2))


# ---------------------------------------------------------------------------
# Dataset compartilhado de experiencias e tabela compartilhada de lances
# ---------------------------------------------------------------------------
class Experiencias:
    """posicao (apos o lance) + de quem foi o lance + resultado esperado (-1..+1)."""

    def __init__(self):
        self.epd, self.cor, self.alvo, self.jogo = [], [], [], []
        self.X, self.F = [], []

    def __len__(self):
        return len(self.alvo)

    def adicionar(self, epd, cor, alvo, x, f, jogo):
        self.epd.append(epd)
        self.cor.append(cor)
        self.alvo.append(alvo)
        self.jogo.append(jogo)
        self.X.append(x)
        self.F.append(f)
        excesso = len(self.alvo) - BUFFER_MAX
        if excesso > 0:
            for lista in (self.epd, self.cor, self.alvo, self.jogo, self.X, self.F):
                del lista[:excesso]

    def lote(self, n):
        idx = random.sample(range(len(self.alvo)), min(n, len(self.alvo)))
        X = np.array([self.X[i] for i in idx])
        F = np.array([self.F[i] for i in idx])
        y = np.array([self.alvo[i] for i in idx])
        return X, F, y

    def carregar(self, linhas):
        for epd, cor, alvo, jogo in linhas:
            try:
                board = chess.Board()
                board.set_epd(epd)
            except ValueError:
                continue
            c = chess.WHITE if cor == "B" else chess.BLACK
            self.adicionar(epd, c, float(alvo), feats_rede(board, c), feats_padroes(board, c),
                           int(jogo or 0))

    def linhas(self):
        for epd, cor, alvo, jogo in zip(self.epd, self.cor, self.alvo, self.jogo):
            yield (epd, "B" if cor == chess.WHITE else "P", round(alvo, 4), jogo)


class Compartilhada:
    """{ posicao: { lance: [vitorias, empates, derrotas] } } do ponto de vista de quem fez o lance."""

    def __init__(self):
        self.dados = {}

    def pontuacao(self, epd, uci):
        r = self.dados.get(epd, {}).get(uci)
        if not r:
            return 0.0
        n = sum(r)
        return ((r[0] - r[2]) / n) * (min(n, 5) / 5)

    def registrar(self, epd, uci, idx):
        """idx: 0 = vitoria, 1 = empate, 2 = derrota."""
        self.dados.setdefault(epd, {}).setdefault(uci, [0, 0, 0])[idx] += 1

    def carregar(self, linhas):
        for epd, uci, v, e, d in linhas:
            self.dados.setdefault(epd, {})[uci] = [int(v or 0), int(e or 0), int(d or 0)]

    def linhas(self):
        for epd, lances in self.dados.items():
            for uci, (v, e, d) in lances.items():
                yield (epd, uci, v, e, d)


# ---------------------------------------------------------------------------
# Jogadores
# ---------------------------------------------------------------------------
class Agente:
    def __init__(self, nome, compartilhada):
        self.nome = nome
        self.comp = compartilhada
        self.historico = []   # (antes, lance_uci, depois, cor, x, f)
        self.ultimo_erro = None

    def novo_jogo(self):
        self.historico = []

    def jogar(self, board):
        antes = board.epd()
        cor = board.turn
        lance = self.escolher(board, antes, cor)
        board.push(lance)
        depois = board.epd()
        x = feats_rede(board, cor)
        f = feats_padroes(board, cor)
        board.pop()
        self.historico.append((antes, lance.uci(), depois, cor, x, f))
        return lance

    def _avaliar_candidatos(self, board, antes, cor, construir):
        """Retorna (lances, matriz_de_features, lance_que_da_mate_ou_None)."""
        lances = list(board.legal_moves)
        linhas = []
        for m in lances:
            board.push(m)
            if board.is_checkmate():
                board.pop()
                return lances, None, m
            linhas.append(construir(board, cor))
            board.pop()
        return lances, np.array(linhas), None

    def bonus_compartilhada(self, antes, lances):
        return PESO_COMPARTILHADA * np.array([self.comp.pontuacao(antes, m.uci()) for m in lances])


class Hero(Agente):
    """Rede neural. Explora com probabilidade epsilon (diminui com o tempo)."""

    def __init__(self, compartilhada):
        super().__init__("Hero", compartilhada)
        self.rede = RedeNeural()
        self.epsilon = 0.2

    def escolher(self, board, antes, cor):
        lances, X, mate = self._avaliar_candidatos(board, antes, cor, feats_rede)
        if mate:
            return mate
        if random.random() < self.epsilon:
            return random.choice(lances)
        valores = self.rede.prever(X) + self.bonus_compartilhada(antes, lances)
        valores = valores + np.random.uniform(-1e-3, 1e-3, len(valores))
        return lances[int(np.argmax(valores))]

    def aprender(self, exp):
        if len(exp) < MIN_TREINO:
            return None
        erros = []
        for _ in range(PASSOS_TREINO):
            X, _, y = exp.lote(LOTE)
            erros.append(self.rede.treinar(X, y))
        self.ultimo_erro = round(float(np.mean(erros[-10:])), 4)
        return self.ultimo_erro


class Boy(Agente):
    """Regressao sobre padroes. Explora com 'temperatura' (softmax) que diminui com o tempo."""

    def __init__(self, compartilhada):
        super().__init__("Boy", compartilhada)
        self.modelo = Regressao()
        self.temperatura = 0.08

    def escolher(self, board, antes, cor):
        lances, F, mate = self._avaliar_candidatos(board, antes, cor, feats_padroes)
        if mate:
            return mate
        valores = self.modelo.prever(F) + self.bonus_compartilhada(antes, lances)
        z = (valores - valores.max()) / max(self.temperatura, 1e-3)
        pesos = np.exp(z)
        return random.choices(lances, weights=pesos / pesos.sum(), k=1)[0]

    def aprender(self, exp):
        if len(exp) < MIN_TREINO:
            return None
        erros = []
        for _ in range(PASSOS_TREINO):
            _, F, y = exp.lote(LOTE)
            erros.append(self.modelo.treinar(F, y))
        self.ultimo_erro = round(float(np.mean(erros[-10:])), 4)
        return self.ultimo_erro


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def _ler_aba(wb, nome):
    if nome not in wb.sheetnames:
        return []
    return [r for r in wb[nome].iter_rows(min_row=2, values_only=True)
            if r and r[0] is not None]


def carregar_tudo():
    comp = Compartilhada()
    hero, boy, exp = Hero(comp), Boy(comp), Experiencias()
    partidas = []
    if os.path.exists(ARQUIVO):
        try:
            wb = load_workbook(ARQUIVO, read_only=True, data_only=True)
            comp.carregar(_ler_aba(wb, "Compartilhada"))
            partidas = [tuple(r)[:11] for r in _ler_aba(wb, "Partidas")]

            w1 = [[float(v) for v in r[1:1 + OCULTA]] for r in _ler_aba(wb, "Hero_Rede")]
            resto = {r[0]: [float(v) for v in r[1:] if v is not None] for r in _ler_aba(wb, "Hero_Resto")}
            if (len(w1) == N_ENTRADA and all(len(l) == OCULTA for l in w1)
                    and len(resto.get("b1", [])) == OCULTA and len(resto.get("W2", [])) == OCULTA
                    and len(resto.get("b2", [])) == 1):
                hero.rede.W1[:] = np.array(w1)
                hero.rede.b1[:] = np.array(resto["b1"])
                hero.rede.W2[:, 0] = np.array(resto["W2"])
                hero.rede.b2[:] = np.array(resto["b2"])

            pesos = [float(r[1]) for r in _ler_aba(wb, "Boy_Modelo")]
            if len(pesos) == N_PADROES:
                boy.modelo.w[:] = np.array(pesos)

            exp.carregar(_ler_aba(wb, "Experiencias"))
            wb.close()
        except Exception as erro:  # arquivo corrompido, formato errado etc.
            backup = ARQUIVO + ".bak"
            try:
                os.replace(ARQUIVO, backup)
            except OSError:
                pass
            print(f"[aviso] Nao consegui ler o Excel ({erro}). Guardei o arquivo antigo em {backup} "
                  f"e vou comecar do zero.")
            comp = Compartilhada()
            hero, boy, exp = Hero(comp), Boy(comp), Experiencias()
            partidas = []
    return hero, boy, comp, exp, partidas


def _aba(wb, nome, cabecalho, linhas, larguras, primeira=False):
    ws = wb.active if primeira else wb.create_sheet()
    ws.title = nome
    ws.append(cabecalho)
    for celula in ws[1]:
        celula.font = Font(bold=True)
    for linha in linhas:
        ws.append(list(linha))
    ws.freeze_panes = "A2"
    for i, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    return ws


def _nomes_entradas():
    nomes = []
    for grupo in range(12):
        dono = "minha" if grupo < 6 else "adversario"
        for casa in range(64):
            nomes.append(f"{dono} {NOMES_PECAS[grupo % 6]} em {chess.square_name(casa)}")
    nomes.append("saldo de material")
    return nomes


def salvar_excel(hero, boy, comp, exp, partidas):
    """Retorna True se salvou. Se o Excel estiver aberto, retorna False."""
    wb = Workbook()

    total = len(partidas)
    v_hero = sum(1 for p in partidas if p[5] == "Hero")
    v_boy = sum(1 for p in partidas if p[5] == "Boy")
    placar = [
        ("Total de partidas", total),
        ("Vitorias do Hero (rede neural)", v_hero),
        ("Vitorias do Boy (regressao)", v_boy),
        ("Empates", total - v_hero - v_boy),
        ("Vitorias do Hero jogando de brancas", sum(1 for p in partidas if p[2] == "Hero" and p[5] == "Hero")),
        ("Vitorias do Boy jogando de brancas", sum(1 for p in partidas if p[2] == "Boy" and p[5] == "Boy")),
        ("Experiencias no dataset de treino", len(exp)),
        ("Posicoes na tabela compartilhada", len(comp.dados)),
        ("Erro atual do Hero (menor = preve melhor)", hero.ultimo_erro),
        ("Erro atual do Boy (menor = preve melhor)", boy.ultimo_erro),
    ]
    ws_placar = _aba(wb, "Placar", ["Item", "Valor"], placar, [44, 12], primeira=True)

    ws_partidas = _aba(
        wb, "Partidas",
        ["N", "Data/Hora", "Brancas", "Pretas", "Resultado", "Vencedor", "Motivo",
         "Meio-lances", "Erro Hero", "Erro Boy", "Lances"],
        partidas, [6, 20, 10, 10, 11, 10, 38, 12, 10, 10, 100])

    if total >= 2:  # curva de aprendizado: o erro deve cair com o tempo
        grafico = LineChart()
        grafico.title = "Curva de aprendizado (erro de previsao)"
        grafico.y_axis.title = "erro"
        grafico.x_axis.title = "partida"
        grafico.add_data(Reference(ws_partidas, min_col=9, max_col=10, min_row=1, max_row=total + 1),
                         titles_from_data=True)
        grafico.set_categories(Reference(ws_partidas, min_col=1, min_row=2, max_row=total + 1))
        grafico.width, grafico.height = 20, 9
        ws_placar.add_chart(grafico, "D2")

    nomes = _nomes_entradas()
    w1 = ([nomes[i]] + [round(float(v), 6) for v in hero.rede.W1[i]] for i in range(N_ENTRADA))
    _aba(wb, "Hero_Rede", ["entrada"] + [f"neuronio {j + 1}" for j in range(OCULTA)],
         w1, [34] + [11] * OCULTA)
    resto = [["b1"] + [round(float(v), 6) for v in hero.rede.b1],
             ["W2"] + [round(float(v), 6) for v in hero.rede.W2[:, 0]],
             ["b2", round(float(hero.rede.b2[0]), 6)]]
    _aba(wb, "Hero_Resto", ["parametro"] + [f"neuronio {j + 1}" for j in range(OCULTA)],
         resto, [12] + [11] * OCULTA)

    _aba(wb, "Boy_Modelo", ["Padrao", "Peso aprendido"],
         [(n, round(float(w), 6)) for n, w in zip(NOMES_PADROES, boy.modelo.w)], [32, 16])
    _aba(wb, "Experiencias", ["Posicao apos o lance (EPD)", "Quem jogou (B/P)", "Alvo (-1 a +1)", "Partida"],
         exp.linhas(), [70, 16, 14, 10])
    _aba(wb, "Compartilhada", ["Posicao (EPD)", "Lance", "Vitorias", "Empates", "Derrotas"],
         comp.linhas(), [70, 10, 10, 10, 10])

    temporario = ARQUIVO.replace(".xlsx", "_tmp.xlsx")
    wb.save(temporario)
    try:
        os.replace(temporario, ARQUIVO)
    except PermissionError:
        print("\n[aviso] Nao consegui salvar: o Excel esta com memoria_xadrez_ml.xlsx aberto. "
              "Feche o arquivo.")
        return False
    return True


class Sair(Exception):
    pass


# ---------------------------------------------------------------------------
# Interface grafica (pygame)
# ---------------------------------------------------------------------------
LARG, ALT = 1120, 720
TAM = 80                 # tamanho de cada casa
BX, BY = 40, 36          # canto do tabuleiro
PX = 700                 # inicio do painel lateral

FUNDO = (28, 30, 36)
CASA_CLARA = (238, 216, 180)
CASA_ESCURA = (176, 128, 92)
COR_HERO = (90, 165, 255)
COR_BOY = (255, 165, 70)
TEXTO = (232, 234, 240)
TEXTO_FRACO = (150, 156, 170)
CARTAO = (46, 50, 60)

GLIFOS = {chess.KING: "\u265a", chess.QUEEN: "\u265b", chess.ROOK: "\u265c",
          chess.BISHOP: "\u265d", chess.KNIGHT: "\u265e", chess.PAWN: "\u265f"}
LETRAS = {chess.KING: "K", chess.QUEEN: "Q", chess.ROOK: "R",
          chess.BISHOP: "B", chess.KNIGHT: "N", chess.PAWN: "P"}


class Contexto:
    """Tudo que a tela precisa saber sobre o jogo atual."""

    def __init__(self, hero, boy, partidas, exp):
        self.hero, self.boy, self.partidas, self.exp = hero, boy, partidas, exp
        self.board = chess.Board()
        self.numero = len(partidas) + 1
        self.brancas = self.pretas = None
        self.lances = []
        self.msg = ""
        self.banner = None
        self.recalcular()

    def recalcular(self):
        self.v_hero = sum(1 for p in self.partidas if p[5] == "Hero")
        self.v_boy = sum(1 for p in self.partidas if p[5] == "Boy")
        self.empates = len(self.partidas) - self.v_hero - self.v_boy


class Interface:
    def __init__(self):
        pygame.init()
        pygame.display.set_caption("Xadrez: Hero (rede neural) x Boy (regressao)")
        self.tela = pygame.display.set_mode((LARG, ALT))
        self.relogio = pygame.time.Clock()
        familia = "segoeui,arial,dejavusans,liberationsans"
        self.f_peq = pygame.font.SysFont(familia, 15)
        self.f_med = pygame.font.SysFont(familia, 18)
        self.f_gra = pygame.font.SysFont(familia, 26, bold=True)
        self.f_tit = pygame.font.SysFont(familia, 36, bold=True)
        self.fonte_glifos = self._achar_fonte_glifos()
        self.glifos = self.fonte_glifos is not None
        self.destaque = pygame.Surface((TAM, TAM), pygame.SRCALPHA)
        self.destaque.fill((246, 246, 105, 110))
        self.alerta = pygame.Surface((TAM, TAM), pygame.SRCALPHA)
        self.alerta.fill((230, 50, 50, 140))
        self.modo = "auto"
        self.pausa = PAUSA_AUTO
        self.espaco = False
        self.botoes = {}
        self._ultimo = 0.0

    # ---------- fontes e pecas ----------
    @staticmethod
    def _achar_fonte_glifos():
        """Procura uma fonte do sistema que tenha as pecas de xadrez (Unicode)."""
        for nome in ("segoeuisymbol", "dejavusans", "freeserif", "symbola",
                     "notosanssymbols2", "arialunicodems", "applesymbols"):
            caminho = pygame.font.match_font(nome)
            if caminho:
                try:
                    fonte = pygame.font.Font(caminho, 60)
                    if all(m is not None for m in fonte.metrics("".join(GLIFOS.values()))):
                        return fonte
                except Exception:
                    pass
        return None

    def texto(self, txt, fonte, cor, x, y, centro=False, direita=False):
        s = fonte.render(txt, True, cor)
        r = s.get_rect()
        if centro:
            r.midtop = (x, y)
        elif direita:
            r.topright = (x, y)
        else:
            r.topleft = (x, y)
        self.tela.blit(s, r)

    def _peca(self, p, cx, cy):
        branca = p.color == chess.WHITE
        if self.glifos and self.fonte_glifos:
            txt = GLIFOS[p.piece_type]
            cor = (252, 252, 252) if branca else (18, 18, 18)
            contorno = (25, 25, 25) if branca else (235, 235, 235)
            borda = self.fonte_glifos.render(txt, True, contorno)
            for dx, dy in ((-2, 0), (2, 0), (0, -2), (0, 2), (-1, -1), (1, 1), (-1, 1), (1, -1)):
                self.tela.blit(borda, borda.get_rect(center=(cx + dx, cy + dy)))
            frente = self.fonte_glifos.render(txt, True, cor)
            self.tela.blit(frente, frente.get_rect(center=(cx, cy)))
        else:
            pygame.draw.circle(self.tela, (250, 250, 250) if branca else (35, 35, 35), (cx, cy), 31)
            pygame.draw.circle(self.tela, (30, 30, 30) if branca else (230, 230, 230), (cx, cy), 31, 3)
            s = self.f_gra.render(LETRAS[p.piece_type], True,
                                  (30, 30, 30) if branca else (240, 240, 240))
            self.tela.blit(s, s.get_rect(center=(cx, cy)))

    # ---------- desenho ----------
    def desenhar(self, ctx, force=False):
        agora = time.time()
        if not force and agora - self._ultimo < 0.1:
            return
        self._ultimo = agora
        self.botoes = {}
        self.tela.fill(FUNDO)
        self._tabuleiro(ctx)
        self._painel(ctx)
        if self.modo is None:
            self._tela_inicial()
        pygame.display.flip()

    def _tabuleiro(self, ctx):
        board = ctx.board
        ultimo = board.peek() if board.move_stack else None
        for r in range(8):
            for f in range(8):
                casa = chess.square(f, r)
                x, y = BX + f * TAM, BY + (7 - r) * TAM
                cor = CASA_CLARA if (f + r) % 2 == 1 else CASA_ESCURA
                pygame.draw.rect(self.tela, cor, (x, y, TAM, TAM))
                if ultimo and casa in (ultimo.from_square, ultimo.to_square):
                    self.tela.blit(self.destaque, (x, y))
        if board.is_check():
            rei = board.king(board.turn)
            if rei is not None:
                self.tela.blit(self.alerta, (BX + chess.square_file(rei) * TAM,
                                             BY + (7 - chess.square_rank(rei)) * TAM))
        for casa, p in board.piece_map().items():
            cx = BX + chess.square_file(casa) * TAM + TAM // 2
            cy = BY + (7 - chess.square_rank(casa)) * TAM + TAM // 2
            self._peca(p, cx, cy)
        for i in range(8):
            self.texto("abcdefgh"[i], self.f_peq, TEXTO_FRACO, BX + i * TAM + TAM // 2,
                       BY + 8 * TAM + 4, centro=True)
            self.texto(str(8 - i), self.f_peq, TEXTO_FRACO, BX - 16, BY + i * TAM + TAM // 2 - 9)
        self.texto(ctx.msg or self._ajuda(), self.f_peq, TEXTO, BX, 10)
        if ctx.banner:
            titulo, sub = ctx.banner
            sup = pygame.Surface((560, 110), pygame.SRCALPHA)
            sup.fill((10, 10, 14, 220))
            self.tela.blit(sup, (BX + 40, BY + 265))
            self.texto(titulo, self.f_gra, TEXTO, BX + 320, BY + 285, centro=True)
            self.texto(sub, self.f_med, TEXTO_FRACO, BX + 320, BY + 328, centro=True)

    def _ajuda(self):
        if self.modo == "auto":
            return f"ESPACO pausa  |  +/- velocidade (pausa {self.pausa:.2f}s)  |  G muda as pecas"
        if self.modo == "rapido":
            return "Modo rapido: ESPACO pausa  |  A/L/P muda o modo  |  Q sai"
        if self.modo == "lance":
            return "ESPACO = proximo lance  |  A/R/P muda o modo  |  Q sai"
        if self.modo == "partida":
            return "Jogando a partida inteira... ESPACO vale no fim da partida"
        return ""

    def _painel(self, ctx):
        x = PX
        self.texto("HERO x BOY", self.f_tit, TEXTO, x, 8)
        self.texto("xadrez com Machine Learning", self.f_peq, TEXTO_FRACO, x, 50)
        modo = NOMES_MODO.get(self.modo, "escolha o modo")
        self.texto(f"Partida {ctx.numero}  |  {modo}", self.f_peq, TEXTO, x, 72)

        self._cartao(x, 96, 380, 84, ctx.hero, "Hero", "rede neural 769-32-1", COR_HERO, ctx, ctx.v_hero)
        self._cartao(x, 186, 380, 84, ctx.boy, "Boy", "regressao 14 padroes", COR_BOY, ctx, ctx.v_boy)

        # placar
        total = max(len(ctx.partidas), 1)
        self.texto(f"Empates: {ctx.empates}   |   Partidas: {len(ctx.partidas)}", self.f_peq,
                   TEXTO_FRACO, x + 190, 276, centro=True)
        larg_h = int(380 * ctx.v_hero / total)
        larg_b = int(380 * ctx.v_boy / total)
        pygame.draw.rect(self.tela, (80, 84, 96), (x, 298, 380, 10), border_radius=4)
        if larg_h:
            pygame.draw.rect(self.tela, COR_HERO, (x, 298, larg_h, 10), border_radius=4)
        if larg_b:
            pygame.draw.rect(self.tela, COR_BOY, (x + 380 - larg_b, 298, larg_b, 10), border_radius=4)

        # o que cada modelo "acha" da posicao atual (visao das brancas)
        self.texto("O que cada modelo acha da posicao (+ = brancas melhor)", self.f_peq,
                   TEXTO_FRACO, x, 316)
        aval_h, aval_b = self._avaliacoes(ctx)
        self._barra(x, 340, 380, "Hero", aval_h, COR_HERO)
        self._barra(x, 364, 380, "Boy", aval_b, COR_BOY)

        self._grafico(ctx, x, 392, 380, 108)

        self.texto("Ultimos lances", self.f_peq, TEXTO_FRACO, x, 508)
        for i, lance in enumerate(ctx.lances[-10:]):
            self.texto(lance, self.f_peq, TEXTO, x + (i // 5) * 190, 528 + (i % 5) * 18)

        # botoes
        nomes = (("auto", "A Auto"), ("rapido", "R Rapido"), ("lance", "L Lance"), ("partida", "P Partida"))
        for i, (nome, rotulo) in enumerate(nomes):
            self._botao(nome, pygame.Rect(x + i * 96, 640, 88, 32), rotulo, ativo=(self.modo == nome))
        self._botao("avancar", pygame.Rect(x, 678, 170, 32), "ESPACO avancar")
        self._botao("glifos", pygame.Rect(x + 178, 678, 100, 32), "G pecas")
        self._botao("sair", pygame.Rect(x + 286, 678, 94, 32), "Q sair")

    def _cartao(self, x, y, w, h, agente, nome, sub, cor, ctx, vitorias):
        pygame.draw.rect(self.tela, CARTAO, (x, y, w, h), border_radius=10)
        pygame.draw.rect(self.tela, cor, (x, y + 10, 5, h - 20))
        self.texto(nome, self.f_gra, cor, x + 16, y + 6)
        self.texto(sub, self.f_peq, TEXTO_FRACO, x + 16, y + 38)
        if ctx.brancas is agente:
            cor_p, rotulo = (245, 245, 245), "joga de BRANCAS"
        elif ctx.pretas is agente:
            cor_p, rotulo = (25, 25, 25), "joga de PRETAS"
        else:
            cor_p = rotulo = None
        if cor_p:
            pygame.draw.circle(self.tela, cor_p, (x + 24, y + 66), 7)
            pygame.draw.circle(self.tela, (200, 200, 200), (x + 24, y + 66), 7, 1)
            self.texto(rotulo, self.f_peq, TEXTO, x + 38, y + 58)
        erro = agente.ultimo_erro
        self.texto(f"erro: {erro:.4f}" if erro is not None else "erro: -", self.f_peq, TEXTO, x + 190, y + 58)
        extra = (f"explora {agente.epsilon:.0%}" if isinstance(agente, Hero)
                 else f"temperatura {agente.temperatura:.2f}")
        self.texto(extra, self.f_peq, TEXTO_FRACO, x + 190, y + 38)
        self.texto(str(vitorias), self.f_tit, TEXTO, x + w - 14, y + 8, direita=True)
        self.texto("vitorias", self.f_peq, TEXTO_FRACO, x + w - 14, y + 54, direita=True)

    @staticmethod
    def _avaliacoes(ctx):
        b = ctx.board
        hero = float(ctx.hero.rede.prever(feats_rede(b, chess.WHITE)[None, :])[0])
        boy = float(ctx.boy.modelo.prever(feats_padroes(b, chess.WHITE)[None, :])[0])
        return hero, boy

    def _barra(self, x, y, w, rotulo, valor, cor):
        self.texto(rotulo, self.f_peq, cor, x, y - 1)
        bx, bw = x + 46, w - 46 - 52
        pygame.draw.rect(self.tela, (20, 22, 28), (bx, y, bw, 14), border_radius=4)
        centro = bx + bw // 2
        ponta = int(centro + max(-1.0, min(1.0, valor)) * bw / 2)
        esq, dire = sorted((centro, ponta))
        pygame.draw.rect(self.tela, cor, (esq, y, max(dire - esq, 1), 14))
        pygame.draw.line(self.tela, TEXTO, (centro, y - 2), (centro, y + 16), 1)
        self.texto(f"{valor:+.2f}", self.f_peq, TEXTO, x + w, y - 1, direita=True)

    def _grafico(self, ctx, x, y, w, h):
        """Curva de aprendizado: o erro de previsao de cada modelo ao longo das partidas."""
        pygame.draw.rect(self.tela, CARTAO, (x, y, w, h), border_radius=8)
        self.texto("curva de aprendizado (erro; menor = melhor)", self.f_peq, TEXTO_FRACO, x + 8, y + 4)
        recentes = ctx.partidas[-120:]
        series = ((COR_HERO, [p[8] for p in recentes]), (COR_BOY, [p[9] for p in recentes]))
        valores = [float(v) for _, s in series for v in s if v is not None]
        if len(recentes) < 2 or not valores:
            self.texto("aguardando partidas para treinar...", self.f_peq, TEXTO_FRACO, x + 8, y + 50)
            return
        topo = max(max(valores), 0.05)
        ax, ay, aw, ah = x + 10, y + 28, w - 20, h - 38
        n = len(recentes)
        for cor, s in series:
            pts = [(ax + aw * i / max(n - 1, 1), ay + ah * (1 - min(float(v), topo) / topo))
                   for i, v in enumerate(s) if v is not None]
            if len(pts) >= 2:
                pygame.draw.lines(self.tela, cor, False, pts, 2)
        self.texto(f"{topo:.2f}", self.f_peq, TEXTO_FRACO, x + w - 8, y + 4, direita=True)

    def _botao(self, nome, rect, rotulo, ativo=False, fonte=None):
        self.botoes[nome] = rect
        cor = (70, 110, 200) if ativo else (62, 66, 78)
        if rect.collidepoint(pygame.mouse.get_pos()):
            cor = tuple(min(255, c + 25) for c in cor)
        pygame.draw.rect(self.tela, cor, rect, border_radius=8)
        s = (fonte or self.f_peq).render(rotulo, True, TEXTO)
        self.tela.blit(s, s.get_rect(center=rect.center))

    def _tela_inicial(self):
        sup = pygame.Surface((600, 440), pygame.SRCALPHA)
        sup.fill((10, 10, 14, 235))
        self.tela.blit(sup, (BX + 20, BY + 100))
        self.texto("Escolha o modo de jogo", self.f_gra, TEXTO, BX + 320, BY + 118, centro=True)
        opcoes = (("auto", "A  Automatico", "Hero e Boy jogam sozinhos e voce assiste"),
                  ("rapido", "R  Rapido", "sem pausas: acompanhe o placar e o aprendizado"),
                  ("lance", "L  Espaco a cada lance", "voce aperta ESPACO para cada lance"),
                  ("partida", "P  Espaco a cada partida", "a partida roda inteira, ESPACO inicia a proxima"))
        for i, (nome, titulo, desc) in enumerate(opcoes):
            rect = pygame.Rect(BX + 50, BY + 170 + i * 84, 540, 72)
            self._botao(nome, rect, "", ativo=False)
            self.texto(titulo, self.f_med, TEXTO, rect.x + 16, rect.y + 10)
            self.texto(desc, self.f_peq, TEXTO_FRACO, rect.x + 16, rect.y + 40)

    # ---------- entrada (teclado e mouse) ----------
    def mudar(self, modo):
        if modo != self.modo:
            self.modo = modo

    def _espaco(self):
        if self.modo in ("auto", "rapido"):
            self.modo = "lance"       # pausa
        else:
            self.espaco = True

    def _acao(self, nome):
        if nome in TECLAS_MODO.values():
            self.mudar(nome)
        elif nome == "avancar":
            self._espaco()
        elif nome == "glifos":
            if self.fonte_glifos:
                self.glifos = not self.glifos
        elif nome == "sair":
            raise Sair()

    def eventos(self):
        for ev in pygame.event.get():
            if ev.type == pygame.QUIT:
                raise Sair()
            if ev.type == pygame.KEYDOWN:
                if ev.key in (pygame.K_ESCAPE, pygame.K_q):
                    raise Sair()
                elif ev.key == pygame.K_SPACE:
                    self._espaco()
                elif ev.key == pygame.K_g:
                    self._acao("glifos")
                elif ev.key in (pygame.K_UP, pygame.K_PLUS, pygame.K_EQUALS, pygame.K_KP_PLUS):
                    self.pausa = max(0.03, self.pausa / 1.5)
                elif ev.key in (pygame.K_DOWN, pygame.K_MINUS, pygame.K_KP_MINUS):
                    self.pausa = min(3.0, self.pausa * 1.5)
                else:
                    letra = pygame.key.name(ev.key)
                    if letra in TECLAS_MODO:
                        self.mudar(TECLAS_MODO[letra])
            elif ev.type == pygame.MOUSEBUTTONDOWN and ev.button == 1:
                for nome, rect in list(self.botoes.items()):
                    if rect.collidepoint(ev.pos):
                        self._acao(nome)
                        break

    # ---------- ritmo do jogo ----------
    def escolher_modo(self, ctx):
        self.modo = None
        while self.modo is None:
            self.eventos()
            self.desenhar(ctx, force=True)
            self.relogio.tick(30)

    def esperar(self, ctx, modo_alvo, mensagem):
        """Fica parado ate apertar ESPACO (ou trocar de modo)."""
        self.espaco = False
        ctx.msg = mensagem
        while self.modo == modo_alvo and not self.espaco:
            self.eventos()
            self.desenhar(ctx, force=True)
            self.relogio.tick(30)
        self.espaco = False
        ctx.msg = ""

    def _pausar(self, ctx, segundos):
        fim = time.time() + segundos
        while time.time() < fim and self.modo == "auto":
            self.eventos()
            self.desenhar(ctx)
            self.relogio.tick(60)

    def passo(self, ctx):
        """Chamado depois de cada lance."""
        self.eventos()
        if self.modo == "lance":
            self.esperar(ctx, "lance", "ESPACO = proximo lance")
        elif self.modo == "auto":
            self.desenhar(ctx, force=True)
            self._pausar(ctx, self.pausa)
        else:
            self.desenhar(ctx)

    def fim_partida(self, ctx):
        if self.modo in ("lance", "partida"):
            self.esperar(ctx, self.modo, "Fim da partida. ESPACO = proxima partida")
        elif self.modo == "auto":
            self.desenhar(ctx, force=True)
            self._pausar(ctx, 2.0)

    def fechar(self):
        pygame.quit()


# ---------------------------------------------------------------------------
# Partida
# ---------------------------------------------------------------------------
def resultado_final(board):
    """Retorna (resultado do ponto de vista das brancas, motivo)."""
    out = board.outcome(claim_draw=True)
    if out:
        if out.winner is None:
            return 0, out.termination.name
        return (1 if out.winner else -1), out.termination.name
    dif = material(board, chess.WHITE)
    if dif >= 3:
        return 1, "LIMITE_MATERIAL"
    if dif <= -3:
        return -1, "LIMITE_MATERIAL"
    return 0, "LIMITE_EMPATE"


def jogar_partida(ctx, ui):
    board = chess.Board()
    ctx.board, ctx.lances, ctx.banner = board, [], None
    brancas, pretas = ctx.brancas, ctx.pretas
    brancas.novo_jogo()
    pretas.novo_jogo()
    ui.desenhar(ctx, force=True)
    if ui.modo == "lance":
        ui.esperar(ctx, "lance", "ESPACO = comecar a partida")

    while board.ply() < MAX_LANCES and not board.is_game_over(claim_draw=True):
        agente = brancas if board.turn == chess.WHITE else pretas
        lance = agente.jogar(board)
        n = board.fullmove_number
        san = board.san(lance)
        foi_branca = board.turn == chess.WHITE
        board.push(lance)
        ctx.lances.append(f"{n}. {san}" if foi_branca else f"{n}... {san}")
        ui.passo(ctx)

    res_brancas, motivo = resultado_final(board)
    lances_txt = chess.Board().variation_san(board.move_stack)
    return board, res_brancas, motivo, lances_txt


def main():
    total = int(sys.argv[1]) if len(sys.argv) > 1 and sys.argv[1].isdigit() else None

    print("=" * 52)
    print("  XADREZ: HERO (rede neural) x BOY (regressao)")
    print("=" * 52)
    hero, boy, comp, exp, partidas = carregar_tudo()
    print(f"Excel: {ARQUIVO}")
    print(f"Partidas ja jogadas: {len(partidas)} | experiencias no dataset: {len(exp)} | "
          f"tabela compartilhada: {len(comp.dados)} posicoes")

    ctx = Contexto(hero, boy, partidas, exp)
    ui = Interface()
    feitas = 0

    try:
        ui.escolher_modo(ctx)
        while total is None or feitas < total:
            # exploracao diminui conforme eles ganham experiencia
            fator = 0.99 ** len(partidas)
            hero.epsilon = max(0.03, 0.20 * fator)
            boy.temperatura = max(0.02, 0.08 * fator)

            # alterna quem joga de brancas (le a ultima partida do Excel)
            hero_brancas = partidas[-1][2] != "Hero" if partidas else True
            ctx.brancas, ctx.pretas = (hero, boy) if hero_brancas else (boy, hero)
            brancas, pretas = ctx.brancas, ctx.pretas
            ctx.numero = len(partidas) + 1

            board, rb, motivo, lances_txt = jogar_partida(ctx, ui)

            if rb == 1:
                vencedor, texto = brancas.nome, "1-0"
            elif rb == -1:
                vencedor, texto = pretas.nome, "0-1"
            else:
                vencedor, texto = "Empate", "1/2-1/2"
            motivo_pt = MOTIVOS.get(motivo, motivo.lower())
            ctx.banner = ("Empate" if vencedor == "Empate" else f"{vencedor} venceu!", motivo_pt)
            ctx.msg = "Treinando Hero e Boy com esta partida..."
            ui.desenhar(ctx, force=True)
            ui.eventos()

            # APRENDIZADO -------------------------------------------------
            for agente, res in ((brancas, rb), (pretas, -rb)):
                n = len(agente.historico)
                idx = {1: 0, 0: 1, -1: 2}[res]
                for i, (antes, uci, depois, cor, x, f) in enumerate(agente.historico):
                    alvo = res * GAMMA ** (n - 1 - i)
                    exp.adicionar(depois, cor, alvo, x, f, ctx.numero)   # dataset compartilhado
                    comp.registrar(antes, uci, idx)                     # tabela compartilhada
            erro_hero = hero.aprender(exp)
            erro_boy = boy.aprender(exp)
            # -------------------------------------------------------------

            partidas.append((ctx.numero, datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                             brancas.nome, pretas.nome, texto, vencedor, motivo_pt, board.ply(),
                             erro_hero, erro_boy, lances_txt))
            ctx.recalcular()
            feitas += 1

            resumo = "empate" if vencedor == "Empate" else f"venceu: {vencedor}"
            print(f">>> Partida {ctx.numero}: {brancas.nome}(B) x {pretas.nome}(P)  {texto}  {resumo}  "
                  f"[{motivo_pt} | {board.ply()} meio-lances]  erro Hero: {erro_hero} | erro Boy: {erro_boy}")

            if ui.modo in ("lance", "partida") or feitas % 5 == 0:
                ctx.msg = "Salvando no Excel..."
                ui.desenhar(ctx, force=True)
                salvar_excel(hero, boy, comp, exp, partidas)
            ctx.msg = ""
            ctx.numero = len(partidas) + 1
            ui.fim_partida(ctx)
    except (Sair, KeyboardInterrupt):
        print("\nEncerrando...")
    finally:
        ui.fechar()

    while not salvar_excel(hero, boy, comp, exp, partidas):
        r = input("Feche o Excel e aperte Enter para tentar de novo (ou S para sair sem salvar): ")
        if r.strip().lower() == "s":
            break

    ctx.recalcular()
    print("\n=== PLACAR GERAL ===")
    print(f"Hero: {ctx.v_hero} | Boy: {ctx.v_boy} | Empates: {ctx.empates} | Total: {len(partidas)}")
    print("\nO que o Boy aprendeu (peso de cada padrao):")
    for nome, peso in zip(NOMES_PADROES, boy.modelo.w):
        print(f"  {nome:<28} {peso:+.3f}")
    print(f"\nSalvo em: {ARQUIVO}")


if __name__ == "__main__":
    main()
