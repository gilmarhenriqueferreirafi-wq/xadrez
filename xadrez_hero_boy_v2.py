"""
XADREZ: HERO x BOY  -  v2  (mesmo jeito de aprender, memorias diferentes + painel de pensamento)
==============================================================================================
Duas IAs jogam xadrez numa janela do pygame e aprendem sozinhas com o resultado de cada partida.

  * Hero e Boy sao IGUAIS no jeito de aprender: a mesma rede neural (769 entradas -> 32 neuronios
    -> nota da posicao), o mesmo treino (backpropagation + Adam), a mesma exploracao (sorteio
    proporcional a nota) e eles NASCEM com os mesmos pesos.
  * O que os diferencia e a MEMORIA: cada um tem a propria rede e o proprio dataset privado de
    experiencias (so aprende com as posicoes que ele mesmo jogou).
  * A unica coisa em comum e a TABELA COMPARTILHADA: vitorias/empates/derrotas por lance, que os
    dois alimentam e consultam (da um pequeno bonus na nota do lance).

PAINEL DE PENSAMENTO: antes de cada lance, cada IA mostra os 5 melhores lances considerados,
com a nota da rede, o efeito da tabela compartilhada (V-E-D = vitorias-empates-derrotas daquele
lance), a chance de ser escolhido e se ela "explorou" (jogou um lance que nao era o de maior nota).
As setas no tabuleiro mostram os 3 melhores lances (a mais grossa/forte = o escolhido).

Tudo e salvo no Excel  memoria_xadrez_v2.xlsx  (mesma pasta; o arquivo da versao anterior
fica intacto). Hero e Boy alternam quem joga de brancas a cada partida.

CONTROLES (teclado ou clique nos botoes), valem a qualquer momento:
  A = automatico        R = rapido (sem pausas)
  L = ESPACO joga um lance        P = ESPACO avanca uma partida
  ESPACO (no automatico/rapido) = pausa     + / - = velocidade do automatico
  S = liga/desliga as setas         G = troca o estilo das pecas       Q ou ESC = salva e sai

Instalar:  python -m pip install chess openpyxl numpy pygame-ce
           (no Python 3.14 use  pygame-ce , o  pygame  comum ainda nao tem versao para ele)
Rodar:     python xadrez_hero_boy_v2.py          (ou:  python xadrez_hero_boy_v2.py 50
                                                  para parar sozinho apos 50 partidas)
IMPORTANTE: feche o Excel antes de rodar, senao o arquivo nao consegue ser salvo.
"""
import math
import os
import random
import sys
import time
from datetime import datetime

import chess
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

try:
    import pygame
except ImportError:
    print("Falta o pygame. Instale com:  python -m pip install pygame-ce")
    sys.exit(1)

PASTA = os.path.dirname(os.path.abspath(__file__))
ARQUIVO = os.path.join(PASTA, "memoria_xadrez_v2.xlsx")

MAX_LANCES = 160          # meio-lances; passou disso a partida e encerrada
PAUSA_AUTO = 0.6          # segundos por lance no modo automatico
PESO_COMPARTILHADA = 0.3  # influencia da tabela compartilhada na nota do lance

# --- Machine Learning (igual para Hero e Boy) ---
N_ENTRADA = 769           # 12 tipos de peca x 64 casas + material
OCULTA = 32               # neuronios
GAMMA = 0.97              # desconto: lances proximos do fim valem mais
BUFFER_MAX = 3000         # experiencias guardadas por IA (as mais recentes)
LOTE = 64                 # tamanho do lote de treino
PASSOS_TREINO = 40        # lotes por partida
MIN_TREINO = 64           # so treina quando houver experiencias suficientes
LR = 0.004                # taxa de aprendizado
TEMP_INICIAL = 0.08       # temperatura da exploracao (softmax) no inicio
TEMP_MINIMA = 0.02
DECAIMENTO = 0.99         # a cada partida a temperatura diminui

VALORES = {chess.PAWN: 1, chess.KNIGHT: 3, chess.BISHOP: 3,
           chess.ROOK: 5, chess.QUEEN: 9, chess.KING: 0}

MOTIVOS = {
    "CHECKMATE": "xeque-mate",
    "STALEMATE": "afogamento",
    "INSUFFICIENT_MATERIAL": "material insuficiente",
    "SEVENTYFIVE_MOVES": "regra dos 75 lances",
    "FIFTY_MOVES": "regra dos 50 lances",
    "FIVEFOLD_REPETITION": "repetição (5x)",
    "THREEFOLD_REPETITION": "repetição (3x)",
    "LIMITE_MATERIAL": "limite de lances (vantagem de material)",
    "LIMITE_EMPATE": "limite de lances (igualdade)",
}
NOMES_MODO = {
    "auto": "AUTOMÁTICO",
    "rapido": "RÁPIDO",
    "lance": "ESPAÇO a cada lance",
    "partida": "ESPAÇO a cada partida",
}
TECLAS_MODO = {"a": "auto", "r": "rapido", "l": "lance", "p": "partida"}
NOMES_PECAS = ["peão", "cavalo", "bispo", "torre", "dama", "rei"]


class Sair(Exception):
    pass


# ---------------------------------------------------------------------------
# Tabuleiro -> numeros (o que a rede "enxerga")
# ---------------------------------------------------------------------------
def material(board, cor):
    """Saldo de material do ponto de vista de 'cor'."""
    total = 0
    for peca, valor in VALORES.items():
        total += valor * len(board.pieces(peca, cor))
        total -= valor * len(board.pieces(peca, not cor))
    return total


def feats_rede(board, cor):
    """Entrada da rede: 768 casas (minhas pecas / pecas do adversario) + material.
    O tabuleiro e espelhado para as pretas, entao a rede sempre 've' do ponto de vista de quem joga."""
    x = np.zeros(N_ENTRADA, dtype=np.float32)
    for casa, p in board.piece_map().items():
        if cor == chess.BLACK:
            casa = chess.square_mirror(casa)
        grupo = (0 if p.color == cor else 6) + p.piece_type - 1
        x[grupo * 64 + casa] = 1.0
    x[768] = material(board, cor) / 10.0
    return x


# ---------------------------------------------------------------------------
# Rede neural (so numpy) - a MESMA para Hero e Boy
# ---------------------------------------------------------------------------
class Adam:
    """Otimizador: ajusta os pesos na direcao que reduz o erro."""

    def __init__(self, params, lr):
        self.params = params
        self.lr = lr
        self.m = [np.zeros_like(p) for p in params]
        self.v = [np.zeros_like(p) for p in params]
        self.t = 0

    def passo(self, grads):
        self.t += 1
        for p, g, m, v in zip(self.params, grads, self.m, self.v):
            m *= 0.9
            m += 0.1 * g
            v *= 0.999
            v += 0.001 * g * g
            mh = m / (1 - 0.9 ** self.t)
            vh = v / (1 - 0.999 ** self.t)
            p -= self.lr * mh / (np.sqrt(vh) + 1e-8)


class RedeNeural:
    """769 entradas -> 32 neuronios (tanh) -> 1 saida (tanh) = nota da posicao entre -1 e +1."""

    def __init__(self, pesos=None):
        if pesos is None:
            rng = np.random.default_rng()
            self.W1 = rng.normal(0, 0.05, (N_ENTRADA, OCULTA))
            self.W1[-1, :] = np.abs(rng.normal(1.5, 0.3, OCULTA))  # ponto de partida: material importa
            self.b1 = np.zeros(OCULTA)
            self.W2 = rng.normal(0.05, 0.05, (OCULTA, 1))
            self.b2 = np.zeros(1)
        else:  # copia exata dos pesos (para os dois nascerem iguais)
            self.W1, self.b1, self.W2, self.b2 = [p.copy() for p in pesos]
        self.otimizador = Adam([self.W1, self.b1, self.W2, self.b2], LR)

    def pesos(self):
        return [self.W1, self.b1, self.W2, self.b2]

    def prever(self, X):
        h = np.tanh(X @ self.W1 + self.b1)
        return np.tanh(h @ self.W2 + self.b2)[:, 0]

    def treinar(self, X, y):
        n = len(X)
        h = np.tanh(X @ self.W1 + self.b1)
        saida = np.tanh(h @ self.W2 + self.b2)[:, 0]
        erro = saida - y
        dz2 = (2 * erro / n * (1 - saida ** 2))[:, None]
        dW2 = h.T @ dz2
        db2 = dz2.sum(axis=0)
        dz1 = (dz2 @ self.W2.T) * (1 - h ** 2)
        dW1 = X.T @ dz1
        db1 = dz1.sum(axis=0)
        self.otimizador.passo([dW1, db1, dW2, db2])
        return float(np.mean(erro ** 2))


# ---------------------------------------------------------------------------
# Memorias
# ---------------------------------------------------------------------------
class Experiencias:
    """Memoria PRIVADA de uma IA: posicao (apos o lance dela) + resultado esperado (-1..+1)."""

    def __init__(self):
        self.epd, self.cor, self.alvo, self.jogo, self.X = [], [], [], [], []

    def __len__(self):
        return len(self.alvo)

    def adicionar(self, epd, cor, alvo, x, jogo):
        self.epd.append(epd)
        self.cor.append(cor)
        self.alvo.append(alvo)
        self.jogo.append(jogo)
        self.X.append(x)
        excesso = len(self.alvo) - BUFFER_MAX
        if excesso > 0:
            for lista in (self.epd, self.cor, self.alvo, self.jogo, self.X):
                del lista[:excesso]

    def lote(self, n):
        idx = random.sample(range(len(self.alvo)), min(n, len(self.alvo)))
        return np.array([self.X[i] for i in idx]), np.array([self.alvo[i] for i in idx])

    def carregar(self, linhas):
        for epd, cor, alvo, jogo in linhas:
            try:
                board = chess.Board()
                board.set_epd(epd)
            except ValueError:
                continue
            c = chess.WHITE if cor == "B" else chess.BLACK
            self.adicionar(epd, c, float(alvo), feats_rede(board, c), int(jogo or 0))

    def linhas(self):
        for epd, cor, alvo, jogo in zip(self.epd, self.cor, self.alvo, self.jogo):
            yield (epd, "B" if cor == chess.WHITE else "P", round(alvo, 4), jogo)


class Compartilhada:
    """Conhecimento PUBLICO: { posicao: { lance: [vitorias, empates, derrotas] } }
    (do ponto de vista de quem fez o lance). Os dois alimentam e consultam."""

    def __init__(self):
        self.dados = {}

    def pontuacao(self, epd, uci):
        r = self.dados.get(epd, {}).get(uci)
        if not r:
            return 0.0
        n = sum(r)
        return ((r[0] - r[2]) / n) * (min(n, 5) / 5)

    def registrar(self, epd, uci, idx):
        """idx: 0 = vitoria, 1 = empate, 2 = derrota."""
        self.dados.setdefault(epd, {}).setdefault(uci, [0, 0, 0])[idx] += 1

    def carregar(self, linhas):
        for epd, uci, v, e, d in linhas:
            self.dados.setdefault(epd, {})[uci] = [int(v or 0), int(e or 0), int(d or 0)]

    def linhas(self):
        for epd, lances in self.dados.items():
            for uci, (v, e, d) in lances.items():
                yield (epd, uci, v, e, d)


# ---------------------------------------------------------------------------
# Jogador (Hero e Boy sao a MESMA classe: so a memoria muda)
# ---------------------------------------------------------------------------
class Jogador:
    def __init__(self, nome, compartilhada, rede=None):
        self.nome = nome
        self.comp = compartilhada
        self.rede = rede or RedeNeural()
        self.exp = Experiencias()          # memoria privada
        self.temperatura = TEMP_INICIAL
        self.historico = []                # (antes, lance_uci, depois, cor, x)
        self.ultimo_erro = None
        self.pensamento = None
        self._antes = None

    def novo_jogo(self):
        self.historico = []
        self.pensamento = None

    def pensar(self, board):
        """Decide o lance e guarda o 'pensamento' (para o painel). Nao joga ainda."""
        antes = board.epd()
        self._antes = antes
        cor = board.turn
        lances = list(board.legal_moves)

        linhas = []
        for m in lances:
            board.push(m)
            if board.is_checkmate():
                board.pop()
                self.pensamento = {"mate": board.san(m), "mate_uci": m.uci(), "opcoes": [],
                                   "explorou": False, "n": len(lances)}
                return m
            linhas.append(feats_rede(board, cor))
            board.pop()

        notas = self.rede.prever(np.array(linhas))
        bonus = PESO_COMPARTILHADA * np.array([self.comp.pontuacao(antes, m.uci()) for m in lances])
        total = notas + bonus
        z = (total - total.max()) / max(self.temperatura, 1e-3)
        pesos = np.exp(z)
        chances = pesos / pesos.sum()

        escolhido = random.choices(range(len(lances)), weights=chances, k=1)[0]
        ordem = np.argsort(-total)
        posicao = {int(idx): pos + 1 for pos, idx in enumerate(ordem)}
        top = [int(i) for i in ordem[:5]]
        if escolhido not in top:
            top.append(escolhido)

        opcoes = []
        for idx in top:
            m = lances[idx]
            reg = self.comp.dados.get(antes, {}).get(m.uci())
            opcoes.append({"san": board.san(m), "uci": m.uci(), "nota": float(notas[idx]),
                           "bonus": float(bonus[idx]), "chance": float(chances[idx]),
                           "ordem": posicao[idx], "registro": tuple(reg) if reg else None,
                           "escolhido": idx == escolhido})
        self.pensamento = {"mate": None, "mate_uci": None, "opcoes": opcoes,
                           "explorou": escolhido != int(ordem[0]), "n": len(lances)}
        return lances[escolhido]

    def confirmar(self, board, lance):
        """Joga o lance no tabuleiro e guarda a experiencia."""
        cor = board.turn
        board.push(lance)
        self.historico.append((self._antes, lance.uci(), board.epd(), cor, feats_rede(board, cor)))

    def registrar_partida(self, resultado, numero):
        """Guarda na memoria privada o que aconteceu (+1 vitoria, 0 empate, -1 derrota)."""
        n = len(self.historico)
        for i, (_, _, depois, cor, x) in enumerate(self.historico):
            self.exp.adicionar(depois, cor, resultado * GAMMA ** (n - 1 - i), x, numero)

    def aprender(self):
        """Treina a rede em lotes sorteados da memoria privada. Retorna o erro de previsao."""
        if len(self.exp) < MIN_TREINO:
            return None
        erros = [self.rede.treinar(*self.exp.lote(LOTE)) for _ in range(PASSOS_TREINO)]
        self.ultimo_erro = round(float(np.mean(erros[-10:])), 4)
        return self.ultimo_erro


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def _ler_aba(wb, nome):
    if nome not in wb.sheetnames:
        return []
    return [r for r in wb[nome].iter_rows(min_row=2, values_only=True)
            if r and r[0] is not None]


def _carregar_jogador(wb, jogador):
    nome = jogador.nome
    w1 = [[float(v) for v in r[1:1 + OCULTA]] for r in _ler_aba(wb, f"{nome}_Rede")]
    resto = {r[0]: [float(v) for v in r[1:] if v is not None] for r in _ler_aba(wb, f"{nome}_Resto")}
    if (len(w1) == N_ENTRADA and all(len(l) == OCULTA for l in w1)
            and len(resto.get("b1", [])) == OCULTA and len(resto.get("W2", [])) == OCULTA
            and len(resto.get("b2", [])) == 1):
        jogador.rede.W1[:] = np.array(w1)
        jogador.rede.b1[:] = np.array(resto["b1"])
        jogador.rede.W2[:, 0] = np.array(resto["W2"])
        jogador.rede.b2[:] = np.array(resto["b2"])
    jogador.exp.carregar(_ler_aba(wb, f"{nome}_Experiencias"))


def novos_jogadores(comp):
    """Hero e Boy nascem com os MESMOS pesos."""
    hero = Jogador("Hero", comp)
    boy = Jogador("Boy", comp, RedeNeural(hero.rede.pesos()))
    return hero, boy


def carregar_tudo():
    comp = Compartilhada()
    hero, boy = novos_jogadores(comp)
    partidas = []
    if os.path.exists(ARQUIVO):
        try:
            wb = load_workbook(ARQUIVO, read_only=True, data_only=True)
            comp.carregar(_ler_aba(wb, "Compartilhada"))
            partidas = [tuple(r)[:11] for r in _ler_aba(wb, "Partidas")]
            _carregar_jogador(wb, hero)
            _carregar_jogador(wb, boy)
            wb.close()
        except Exception as erro:  # arquivo corrompido, formato errado etc.
            backup = ARQUIVO + ".bak"
            try:
                os.replace(ARQUIVO, backup)
            except OSError:
                pass
            print(f"[aviso] Nao consegui ler o Excel ({erro}). Guardei o arquivo antigo em {backup} "
                  f"e vou comecar do zero.")
            comp = Compartilhada()
            hero, boy = novos_jogadores(comp)
            partidas = []
    return hero, boy, comp, partidas


def _aba(wb, nome, cabecalho, linhas, larguras, primeira=False):
    ws = wb.active if primeira else wb.create_sheet()
    ws.title = nome
    ws.append(cabecalho)
    for celula in ws[1]:
        celula.font = Font(bold=True)
    for linha in linhas:
        ws.append(list(linha))
    ws.freeze_panes = "A2"
    for i, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    return ws


def _nomes_entradas():
    nomes = []
    for grupo in range(12):
        dono = "minha" if grupo < 6 else "adversario"
        for casa in range(64):
            nomes.append(f"{dono} {NOMES_PECAS[grupo % 6]} em {chess.square_name(casa)}")
    nomes.append("saldo de material")
    return nomes


def _salvar_jogador(wb, jogador, nomes):
    nome = jogador.nome
    rede = jogador.rede
    w1 = ([nomes[i]] + [round(float(v), 6) for v in rede.W1[i]] for i in range(N_ENTRADA))
    _aba(wb, f"{nome}_Rede", ["entrada"] + [f"neuronio {j + 1}" for j in range(OCULTA)],
         w1, [34] + [11] * OCULTA)
    resto = [["b1"] + [round(float(v), 6) for v in rede.b1],
             ["W2"] + [round(float(v), 6) for v in rede.W2[:, 0]],
             ["b2", round(float(rede.b2[0]), 6)]]
    _aba(wb, f"{nome}_Resto", ["parametro"] + [f"neuronio {j + 1}" for j in range(OCULTA)],
         resto, [12] + [11] * OCULTA)
    _aba(wb, f"{nome}_Experiencias",
         ["Posicao apos o lance (EPD)", "Quem jogou (B/P)", "Alvo (-1 a +1)", "Partida"],
         jogador.exp.linhas(), [70, 16, 14, 10])


def salvar_excel(hero, boy, comp, partidas):
    """Retorna True se salvou. Se o Excel estiver aberto, retorna False."""
    wb = Workbook()
    total = len(partidas)
    v_hero = sum(1 for p in partidas if p[5] == "Hero")
    v_boy = sum(1 for p in partidas if p[5] == "Boy")
    placar = [
        ("Total de partidas", total),
        ("Vitorias do Hero", v_hero),
        ("Vitorias do Boy", v_boy),
        ("Empates", total - v_hero - v_boy),
        ("Vitorias do Hero jogando de brancas", sum(1 for p in partidas if p[2] == "Hero" and p[5] == "Hero")),
        ("Vitorias do Boy jogando de brancas", sum(1 for p in partidas if p[2] == "Boy" and p[5] == "Boy")),
        ("Memoria privada do Hero (experiencias)", len(hero.exp)),
        ("Memoria privada do Boy (experiencias)", len(boy.exp)),
        ("Posicoes na tabela compartilhada", len(comp.dados)),
        ("Erro atual do Hero (menor = preve melhor)", hero.ultimo_erro),
        ("Erro atual do Boy (menor = preve melhor)", boy.ultimo_erro),
    ]
    ws_placar = _aba(wb, "Placar", ["Item", "Valor"], placar, [44, 12], primeira=True)

    ws_partidas = _aba(
        wb, "Partidas",
        ["N", "Data/Hora", "Brancas", "Pretas", "Resultado", "Vencedor", "Motivo",
         "Meio-lances", "Erro Hero", "Erro Boy", "Lances"],
        partidas, [6, 20, 10, 10, 11, 10, 38, 12, 10, 10, 100])

    if total >= 2:  # curva de aprendizado: o erro deve cair com o tempo
        grafico = LineChart()
        grafico.title = "Curva de aprendizado (erro de previsao)"
        grafico.y_axis.title = "erro"
        grafico.x_axis.title = "partida"
        grafico.add_data(Reference(ws_partidas, min_col=9, max_col=10, min_row=1, max_row=total + 1),
                         titles_from_data=True)
        grafico.set_categories(Reference(ws_partidas, min_col=1, min_row=2, max_row=total + 1))
        grafico.width, grafico.height = 20, 9
        ws_placar.add_chart(grafico, "D2")

    nomes = _nomes_entradas()
    _salvar_jogador(wb, hero, nomes)
    _salvar_jogador(wb, boy, nomes)
    _aba(wb, "Compartilhada", ["Posicao (EPD)", "Lance", "Vitorias", "Empates", "Derrotas"],
         comp.linhas(), [70, 10, 10, 10, 10])

    temporario = ARQUIVO.replace(".xlsx", "_tmp.xlsx")
    wb.save(temporario)
    try:
        os.replace(temporario, ARQUIVO)
    except PermissionError:
        print("\n[aviso] Nao consegui salvar: o Excel esta com o arquivo aberto. Feche o arquivo.")
        return False
    return True


# ---------------------------------------------------------------------------
# Interface grafica (pygame)
# ---------------------------------------------------------------------------
LARG, ALT = 1120, 720
TAM = 80                 # tamanho de cada casa
BX, BY = 40, 36          # canto do tabuleiro
PX = 700                 # inicio do painel lateral

FUNDO = (28, 30, 36)
CASA_CLARA = (238, 216, 180)
CASA_ESCURA = (176, 128, 92)
COR_HERO = (90, 165, 255)
COR_BOY = (255, 165, 70)
COR_MATE = (90, 220, 120)
COR_EXPLOROU = (255, 214, 90)
TEXTO = (232, 234, 240)
TEXTO_FRACO = (150, 156, 170)
CARTAO = (46, 50, 60)

GLIFOS = {chess.KING: "\u265a", chess.QUEEN: "\u265b", chess.ROOK: "\u265c",
          chess.BISHOP: "\u265d", chess.KNIGHT: "\u265e", chess.PAWN: "\u265f"}
LETRAS = {chess.KING: "K", chess.QUEEN: "Q", chess.ROOK: "R",
          chess.BISHOP: "B", chess.KNIGHT: "N", chess.PAWN: "P"}


class Contexto:
    """Tudo que a tela precisa saber sobre o jogo atual."""

    def __init__(self, hero, boy, partidas):
        self.hero, self.boy, self.partidas = hero, boy, partidas
        self.board = chess.Board()
        self.numero = len(partidas) + 1
        self.brancas = self.pretas = None
        self.pensando = None       # quem esta decidindo o lance agora
        self.ultimo = ""           # texto do ultimo lance jogado
        self.msg = ""
        self.banner = None
        self.recalcular()

    def recalcular(self):
        self.v_hero = sum(1 for p in self.partidas if p[5] == "Hero")
        self.v_boy = sum(1 for p in self.partidas if p[5] == "Boy")
        self.empates = len(self.partidas) - self.v_hero - self.v_boy


class Interface:
    def __init__(self):
        pygame.init()
        pygame.display.set_caption("Xadrez: Hero x Boy (mesma rede neural, memorias diferentes)")
        self.tela = pygame.display.set_mode((LARG, ALT))
        self.relogio = pygame.time.Clock()
        familia = "segoeui,arial,dejavusans,liberationsans"
        self.f_peq = pygame.font.SysFont(familia, 15)
        self.f_med = pygame.font.SysFont(familia, 18)
        self.f_gra = pygame.font.SysFont(familia, 26, bold=True)
        self.f_tit = pygame.font.SysFont(familia, 36, bold=True)
        self.fonte_glifos = self._achar_fonte_glifos()
        self.glifos = self.fonte_glifos is not None
        self.setas = True
        self.destaque = pygame.Surface((TAM, TAM), pygame.SRCALPHA)
        self.destaque.fill((246, 246, 105, 110))
        self.alerta = pygame.Surface((TAM, TAM), pygame.SRCALPHA)
        self.alerta.fill((230, 50, 50, 140))
        self.modo = "auto"
        self.pausa = PAUSA_AUTO
        self.espaco = False
        self.botoes = {}
        self._ultimo = 0.0

    # ---------- fontes e pecas ----------
    @staticmethod
    def _achar_fonte_glifos():
        """Procura uma fonte do sistema que tenha as pecas de xadrez (Unicode)."""
        for nome in ("segoeuisymbol", "dejavusans", "freeserif", "symbola",
                     "notosanssymbols2", "arialunicodems", "applesymbols"):
            caminho = pygame.font.match_font(nome)
            if caminho:
                try:
                    fonte = pygame.font.Font(caminho, 60)
                    if all(m is not None for m in fonte.metrics("".join(GLIFOS.values()))):
                        return fonte
                except Exception:
                    pass
        return None

    def texto(self, txt, fonte, cor, x, y, centro=False, direita=False):
        s = fonte.render(txt, True, cor)
        r = s.get_rect()
        if centro:
            r.midtop = (x, y)
        elif direita:
            r.topright = (x, y)
        else:
            r.topleft = (x, y)
        self.tela.blit(s, r)

    def _peca(self, p, cx, cy):
        branca = p.color == chess.WHITE
        if self.glifos and self.fonte_glifos:
            txt = GLIFOS[p.piece_type]
            cor = (252, 252, 252) if branca else (18, 18, 18)
            contorno = (25, 25, 25) if branca else (235, 235, 235)
            borda = self.fonte_glifos.render(txt, True, contorno)
            for dx, dy in ((-2, 0), (2, 0), (0, -2), (0, 2), (-1, -1), (1, 1), (-1, 1), (1, -1)):
                self.tela.blit(borda, borda.get_rect(center=(cx + dx, cy + dy)))
            frente = self.fonte_glifos.render(txt, True, cor)
            self.tela.blit(frente, frente.get_rect(center=(cx, cy)))
        else:
            pygame.draw.circle(self.tela, (250, 250, 250) if branca else (35, 35, 35), (cx, cy), 31)
            pygame.draw.circle(self.tela, (30, 30, 30) if branca else (230, 230, 230), (cx, cy), 31, 3)
            s = self.f_gra.render(LETRAS[p.piece_type], True,
                                  (30, 30, 30) if branca else (240, 240, 240))
            self.tela.blit(s, s.get_rect(center=(cx, cy)))

    # ---------- desenho ----------
    def desenhar(self, ctx, force=False):
        agora = time.time()
        if not force and agora - self._ultimo < 0.1:
            return
        self._ultimo = agora
        self.botoes = {}
        self.tela.fill(FUNDO)
        self._tabuleiro(ctx)
        self._painel(ctx)
        if self.modo is None:
            self._tela_inicial()
        pygame.display.flip()

    @staticmethod
    def _centro_casa(casa):
        return (chess.square_file(casa) * TAM + TAM // 2,
                (7 - chess.square_rank(casa)) * TAM + TAM // 2)

    @staticmethod
    def _seta(camada, de, para, cor, largura):
        dx, dy = para[0] - de[0], para[1] - de[1]
        dist = math.hypot(dx, dy)
        if dist < 1:
            return
        ux, uy = dx / dist, dy / dist
        cabeca = 26
        ponta = (para[0] - ux * 6, para[1] - uy * 6)
        base = (ponta[0] - ux * cabeca, ponta[1] - uy * cabeca)
        pygame.draw.line(camada, cor, de, (base[0] + ux * 4, base[1] + uy * 4), largura)
        px, py = -uy, ux
        pygame.draw.polygon(camada, cor, [ponta,
                                          (base[0] + px * cabeca * 0.6, base[1] + py * cabeca * 0.6),
                                          (base[0] - px * cabeca * 0.6, base[1] - py * cabeca * 0.6)])

    def _setas_pensamento(self, ctx):
        ag = ctx.pensando
        if not (self.setas and ag and ag.pensamento):
            return
        p = ag.pensamento
        cor = COR_HERO if ag is ctx.hero else COR_BOY
        camada = pygame.Surface((8 * TAM, 8 * TAM), pygame.SRCALPHA)
        if p["mate"]:
            m = chess.Move.from_uci(p["mate_uci"])
            self._seta(camada, self._centro_casa(m.from_square), self._centro_casa(m.to_square),
                       COR_MATE + (235,), 14)
        else:
            sel = [o for o in p["opcoes"] if o["escolhido"] or o["ordem"] <= 3]
            for o in sorted(sel, key=lambda o: o["escolhido"]):   # o escolhido fica por cima
                m = chess.Move.from_uci(o["uci"])
                alfa = 240 if o["escolhido"] else 115
                largura = int(6 + 14 * o["chance"]) + (4 if o["escolhido"] else 0)
                self._seta(camada, self._centro_casa(m.from_square), self._centro_casa(m.to_square),
                           cor + (alfa,), largura)
        self.tela.blit(camada, (BX, BY))

    def _tabuleiro(self, ctx):
        board = ctx.board
        ultimo = board.peek() if board.move_stack else None
        for r in range(8):
            for f in range(8):
                casa = chess.square(f, r)
                x, y = BX + f * TAM, BY + (7 - r) * TAM
                cor = CASA_CLARA if (f + r) % 2 == 1 else CASA_ESCURA
                pygame.draw.rect(self.tela, cor, (x, y, TAM, TAM))
                if ultimo and casa in (ultimo.from_square, ultimo.to_square):
                    self.tela.blit(self.destaque, (x, y))
        if board.is_check():
            rei = board.king(board.turn)
            if rei is not None:
                self.tela.blit(self.alerta, (BX + chess.square_file(rei) * TAM,
                                             BY + (7 - chess.square_rank(rei)) * TAM))
        for casa, p in board.piece_map().items():
            cx = BX + chess.square_file(casa) * TAM + TAM // 2
            cy = BY + (7 - chess.square_rank(casa)) * TAM + TAM // 2
            self._peca(p, cx, cy)
        self._setas_pensamento(ctx)
        for i in range(8):
            self.texto("abcdefgh"[i], self.f_peq, TEXTO_FRACO, BX + i * TAM + TAM // 2,
                       BY + 8 * TAM + 4, centro=True)
            self.texto(str(8 - i), self.f_peq, TEXTO_FRACO, BX - 16, BY + i * TAM + TAM // 2 - 9)
        self.texto(ctx.msg or self._ajuda(), self.f_peq, TEXTO, BX, 10)
        if ctx.ultimo:
            self.texto(ctx.ultimo, self.f_peq, TEXTO_FRACO, BX + 8 * TAM, 10, direita=True)
        if ctx.banner:
            titulo, sub = ctx.banner
            sup = pygame.Surface((560, 110), pygame.SRCALPHA)
            sup.fill((10, 10, 14, 220))
            self.tela.blit(sup, (BX + 40, BY + 265))
            self.texto(titulo, self.f_gra, TEXTO, BX + 320, BY + 285, centro=True)
            self.texto(sub, self.f_med, TEXTO_FRACO, BX + 320, BY + 328, centro=True)

    def _ajuda(self):
        if self.modo == "auto":
            return f"ESPAÇO pausa  |  +/- velocidade ({self.pausa:.2f}s)  |  S setas  |  G peças"
        if self.modo == "rapido":
            return "Modo rápido  |  ESPAÇO pausa  |  A/L/P muda o modo  |  Q sai"
        if self.modo == "lance":
            return "ESPAÇO = jogar este lance  |  A/R/P muda o modo  |  S setas  |  Q sai"
        if self.modo == "partida":
            return "Jogando a partida inteira...  ESPAÇO vale no fim da partida"
        return ""

    def _painel(self, ctx):
        x = PX
        self.texto("HERO x BOY", self.f_tit, TEXTO, x, 8)
        self.texto("mesma rede neural, memórias diferentes", self.f_peq, TEXTO_FRACO, x, 50)
        modo = NOMES_MODO.get(self.modo, "escolha o modo")
        self.texto(f"Partida {ctx.numero}  |  {modo}", self.f_peq, TEXTO, x, 72)

        self._cartao(x, 96, 380, 212, ctx.hero, COR_HERO, ctx, ctx.v_hero)
        self._cartao(x, 314, 380, 212, ctx.boy, COR_BOY, ctx, ctx.v_boy)

        # placar
        total = max(len(ctx.partidas), 1)
        self.texto(f"Empates: {ctx.empates}   |   Partidas: {len(ctx.partidas)}", self.f_peq,
                   TEXTO_FRACO, x + 190, 532, centro=True)
        larg_h = int(380 * ctx.v_hero / total)
        larg_b = int(380 * ctx.v_boy / total)
        pygame.draw.rect(self.tela, (80, 84, 96), (x, 552, 380, 10), border_radius=4)
        if larg_h:
            pygame.draw.rect(self.tela, COR_HERO, (x, 552, larg_h, 10), border_radius=4)
        if larg_b:
            pygame.draw.rect(self.tela, COR_BOY, (x + 380 - larg_b, 552, larg_b, 10), border_radius=4)

        self._grafico(ctx, x, 568, 380, 64)

        nomes = (("auto", "A Auto"), ("rapido", "R Rápido"), ("lance", "L Lance"), ("partida", "P Partida"))
        for i, (nome, rotulo) in enumerate(nomes):
            self._botao(nome, pygame.Rect(x + i * 96, 640, 88, 32), rotulo, ativo=(self.modo == nome))
        self._botao("avancar", pygame.Rect(x, 678, 112, 32), "ESPAÇO")
        self._botao("setas", pygame.Rect(x + 120, 678, 84, 32), "S setas", ativo=self.setas)
        self._botao("glifos", pygame.Rect(x + 212, 678, 84, 32), "G peças")
        self._botao("sair", pygame.Rect(x + 304, 678, 76, 32), "Q sair")

    def _cartao(self, x, y, w, h, ag, cor, ctx, vitorias):
        pygame.draw.rect(self.tela, CARTAO, (x, y, w, h), border_radius=10)
        if ctx.pensando is ag:
            pygame.draw.rect(self.tela, cor, (x, y, w, h), 2, border_radius=10)
        pygame.draw.rect(self.tela, cor, (x, y + 10, 5, 62))
        self.texto(ag.nome, self.f_gra, cor, x + 16, y + 4)
        self.texto("rede neural 769-32-1" + ("  -  pensando..." if ctx.pensando is ag else ""),
                   self.f_peq, TEXTO_FRACO, x + 84, y + 12)
        self.texto(str(vitorias), self.f_tit, TEXTO, x + w - 14, y + 2, direita=True)
        self.texto("vitórias", self.f_peq, TEXTO_FRACO, x + w - 14, y + 40, direita=True)

        if ctx.brancas is ag:
            cor_p, rotulo = (245, 245, 245), "joga de BRANCAS"
        elif ctx.pretas is ag:
            cor_p, rotulo = (25, 25, 25), "joga de PRETAS"
        else:
            cor_p = rotulo = None
        if cor_p:
            pygame.draw.circle(self.tela, cor_p, (x + 24, y + 48), 7)
            pygame.draw.circle(self.tela, (200, 200, 200), (x + 24, y + 48), 7, 1)
            self.texto(rotulo, self.f_peq, TEXTO, x + 38, y + 40)
        erro = ag.ultimo_erro
        self.texto(f"erro: {erro:.4f}" if erro is not None else "erro: -", self.f_peq, TEXTO, x + 170, y + 40)
        self.texto(f"temperatura {ag.temperatura:.2f}   |   memória privada: {len(ag.exp)}",
                   self.f_peq, TEXTO_FRACO, x + 16, y + 58)
        self._pensamento(ag, cor, x, y + 80, w)

    def _pensamento(self, ag, cor, x, y, w):
        self.texto("lance", self.f_peq, TEXTO_FRACO, x + 22, y)
        self.texto("nota", self.f_peq, TEXTO_FRACO, x + 92, y)
        self.texto("tabela (V-E-D)", self.f_peq, TEXTO_FRACO, x + 160, y)
        self.texto("chance", self.f_peq, TEXTO_FRACO, x + w - 14, y, direita=True)
        p = ag.pensamento
        if not p:
            self.texto("ainda não jogou nesta partida", self.f_peq, TEXTO_FRACO, x + 14, y + 22)
            return
        if p["mate"]:
            self.texto(f"Viu xeque-mate em 1:  {p['mate']}", self.f_med, COR_MATE, x + 14, y + 22)
            return
        fundo_barra = tuple(int(c * 0.30 + k * 0.70) for c, k in zip(cor, CARTAO))
        fundo_barra_esc = tuple(int(c * 0.55 + k * 0.45) for c, k in zip(cor, CARTAO))
        for i, o in enumerate(p["opcoes"]):
            ly = y + 18 + i * 17
            larg = int((w - 28) * o["chance"])
            if larg > 0:
                pygame.draw.rect(self.tela, fundo_barra_esc if o["escolhido"] else fundo_barra,
                                 (x + 14, ly, larg, 15), border_radius=3)
            cor_txt = TEXTO if o["escolhido"] else TEXTO_FRACO
            if o["escolhido"]:
                pygame.draw.polygon(self.tela, cor, [(x + 4, ly + 2), (x + 4, ly + 13), (x + 13, ly + 7)])
            nome = o["san"] + (f" (#{o['ordem']})" if o["ordem"] > 5 else "")
            self.texto(nome, self.f_peq, cor_txt, x + 22, ly)
            self.texto(f"{o['nota']:+.3f}", self.f_peq, cor_txt, x + 92, ly)
            reg = o["registro"]
            self.texto(f"{o['bonus']:+.2f} ({reg[0]}-{reg[1]}-{reg[2]})" if reg else "-",
                       self.f_peq, cor_txt, x + 160, ly)
            explorou = o["escolhido"] and p["explorou"]
            self.texto(f"{o['chance']:.0%}" + ("  explorou" if explorou else ""), self.f_peq,
                       COR_EXPLOROU if explorou else cor_txt, x + w - 14, ly, direita=True)

    def _grafico(self, ctx, x, y, w, h):
        """Curva de aprendizado: o erro de previsao de cada IA ao longo das partidas."""
        pygame.draw.rect(self.tela, CARTAO, (x, y, w, h), border_radius=8)
        self.texto("erro de previsão (menor = melhor)", self.f_peq, TEXTO_FRACO, x + 8, y + 3)
        recentes = ctx.partidas[-120:]
        series = ((COR_HERO, [p[8] for p in recentes]), (COR_BOY, [p[9] for p in recentes]))
        valores = [float(v) for _, s in series for v in s if v is not None]
        if len(recentes) < 2 or not valores:
            self.texto("aguardando partidas para treinar...", self.f_peq, TEXTO_FRACO, x + 8, y + 32)
            return
        topo = max(max(valores), 0.05)
        ax, ay, aw, ah = x + 10, y + 24, w - 20, h - 30
        n = len(recentes)
        for cor, s in series:
            pts = [(ax + aw * i / max(n - 1, 1), ay + ah * (1 - min(float(v), topo) / topo))
                   for i, v in enumerate(s) if v is not None]
            if len(pts) >= 2:
                pygame.draw.lines(self.tela, cor, False, pts, 2)
        self.texto(f"{topo:.2f}", self.f_peq, TEXTO_FRACO, x + w - 8, y + 3, direita=True)

    def _botao(self, nome, rect, rotulo, ativo=False):
        self.botoes[nome] = rect
        cor = (70, 110, 200) if ativo else (62, 66, 78)
        if rect.collidepoint(pygame.mouse.get_pos()):
            cor = tuple(min(255, c + 25) for c in cor)
        pygame.draw.rect(self.tela, cor, rect, border_radius=8)
        s = self.f_peq.render(rotulo, True, TEXTO)
        self.tela.blit(s, s.get_rect(center=rect.center))

    def _tela_inicial(self):
        sup = pygame.Surface((600, 440), pygame.SRCALPHA)
        sup.fill((10, 10, 14, 235))
        self.tela.blit(sup, (BX + 20, BY + 100))
        self.texto("Escolha o modo de jogo", self.f_gra, TEXTO, BX + 320, BY + 118, centro=True)
        opcoes = (("auto", "A  Automático", "Hero e Boy jogam sozinhos e você assiste os pensamentos"),
                  ("rapido", "R  Rápido", "sem pausas: acompanhe o placar e o aprendizado"),
                  ("lance", "L  ESPAÇO a cada lance", "você aperta ESPAÇO para cada lance (dá tempo de ler o pensamento)"),
                  ("partida", "P  ESPAÇO a cada partida", "a partida roda inteira, ESPAÇO inicia a próxima"))
        for i, (nome, titulo, desc) in enumerate(opcoes):
            rect = pygame.Rect(BX + 50, BY + 170 + i * 84, 540, 72)
            self._botao(nome, rect, "", ativo=False)
            self.texto(titulo, self.f_med, TEXTO, rect.x + 16, rect.y + 10)
            self.texto(desc, self.f_peq, TEXTO_FRACO, rect.x + 16, rect.y + 40)

    # ---------- entrada (teclado e mouse) ----------
    def _espaco(self):
        if self.modo in ("auto", "rapido"):
            self.modo = "lance"       # pausa
        else:
            self.espaco = True

    def _acao(self, nome):
        if nome in TECLAS_MODO.values():
            self.modo = nome
        elif nome == "avancar":
            self._espaco()
        elif nome == "glifos":
            if self.fonte_glifos:
                self.glifos = not self.glifos
        elif nome == "setas":
            self.setas = not self.setas
        elif nome == "sair":
            raise Sair()

    def eventos(self):
        for ev in pygame.event.get():
            if ev.type == pygame.QUIT:
                raise Sair()
            if ev.type == pygame.KEYDOWN:
                if ev.key in (pygame.K_ESCAPE, pygame.K_q):
                    raise Sair()
                elif ev.key == pygame.K_SPACE:
                    self._espaco()
                elif ev.key == pygame.K_g:
                    self._acao("glifos")
                elif ev.key == pygame.K_s:
                    self._acao("setas")
                elif ev.key in (pygame.K_UP, pygame.K_PLUS, pygame.K_EQUALS, pygame.K_KP_PLUS):
                    self.pausa = max(0.03, self.pausa / 1.5)
                elif ev.key in (pygame.K_DOWN, pygame.K_MINUS, pygame.K_KP_MINUS):
                    self.pausa = min(4.0, self.pausa * 1.5)
                else:
                    letra = pygame.key.name(ev.key)
                    if letra in TECLAS_MODO:
                        self.modo = TECLAS_MODO[letra]
            elif ev.type == pygame.MOUSEBUTTONDOWN and ev.button == 1:
                for nome, rect in list(self.botoes.items()):
                    if rect.collidepoint(ev.pos):
                        self._acao(nome)
                        break

    # ---------- ritmo do jogo ----------
    def escolher_modo(self, ctx):
        self.modo = None
        while self.modo is None:
            self.eventos()
            self.desenhar(ctx, force=True)
            self.relogio.tick(30)

    def esperar(self, ctx, modo_alvo, mensagem):
        """Fica parado ate apertar ESPACO (ou trocar de modo)."""
        self.espaco = False
        ctx.msg = mensagem
        while self.modo == modo_alvo and not self.espaco:
            self.eventos()
            self.desenhar(ctx, force=True)
            self.relogio.tick(30)
        self.espaco = False
        ctx.msg = ""

    def _pausar(self, ctx, segundos):
        fim = time.time() + segundos
        while time.time() < fim and self.modo == "auto":
            self.eventos()
            self.desenhar(ctx)
            self.relogio.tick(60)

    def passo(self, ctx):
        """Chamado quando a IA ja decidiu o lance mas ainda nao jogou: mostra o pensamento."""
        self.eventos()
        if self.modo == "lance":
            self.esperar(ctx, "lance", "ESPAÇO = jogar este lance  |  A/R/P muda o modo  |  Q sai")
        elif self.modo == "auto":
            self.desenhar(ctx, force=True)
            self._pausar(ctx, self.pausa)
        else:
            self.desenhar(ctx)

    def fim_partida(self, ctx):
        if self.modo in ("lance", "partida"):
            self.esperar(ctx, self.modo, "Fim da partida.  ESPAÇO = próxima partida")
        elif self.modo == "auto":
            self.desenhar(ctx, force=True)
            self._pausar(ctx, 2.0)

    def fechar(self):
        pygame.quit()


# ---------------------------------------------------------------------------
# Partida
# ---------------------------------------------------------------------------
def resultado_final(board):
    """Retorna (resultado do ponto de vista das brancas, motivo)."""
    out = board.outcome(claim_draw=True)
    if out:
        if out.winner is None:
            return 0, out.termination.name
        return (1 if out.winner else -1), out.termination.name
    dif = material(board, chess.WHITE)
    if dif >= 3:
        return 1, "LIMITE_MATERIAL"
    if dif <= -3:
        return -1, "LIMITE_MATERIAL"
    return 0, "LIMITE_EMPATE"


def jogar_partida(ctx, ui):
    board = chess.Board()
    ctx.board, ctx.banner, ctx.pensando, ctx.ultimo = board, None, None, ""
    brancas, pretas = ctx.brancas, ctx.pretas
    brancas.novo_jogo()
    pretas.novo_jogo()
    ui.desenhar(ctx, force=True)

    while board.ply() < MAX_LANCES and not board.is_game_over(claim_draw=True):
        agente = brancas if board.turn == chess.WHITE else pretas
        lance = agente.pensar(board)          # 1) decide (e guarda o pensamento)
        ctx.pensando = agente
        ui.passo(ctx)                         # 2) mostra o pensamento e espera
        san = board.san(lance)
        n = board.fullmove_number
        foi_branca = board.turn == chess.WHITE
        agente.confirmar(board, lance)        # 3) joga
        ctx.ultimo = f"{agente.nome} jogou {n}. {san}" if foi_branca else f"{agente.nome} jogou {n}... {san}"

    ctx.pensando = None
    res_brancas, motivo = resultado_final(board)
    lances_txt = chess.Board().variation_san(board.move_stack)
    return board, res_brancas, motivo, lances_txt


def main():
    total = int(sys.argv[1]) if len(sys.argv) > 1 and sys.argv[1].isdigit() else None

    print("=" * 56)
    print("  XADREZ: HERO x BOY  (mesma rede neural, memorias diferentes)")
    print("=" * 56)
    hero, boy, comp, partidas = carregar_tudo()
    print(f"Excel: {ARQUIVO}")
    print(f"Partidas ja jogadas: {len(partidas)} | memoria privada Hero: {len(hero.exp)} | "
          f"Boy: {len(boy.exp)} | tabela compartilhada: {len(comp.dados)} posicoes")

    ctx = Contexto(hero, boy, partidas)
    ui = Interface()
    feitas = 0

    try:
        ui.escolher_modo(ctx)
        while total is None or feitas < total:
            # a exploracao diminui conforme eles ganham experiencia (igual para os dois)
            temp = max(TEMP_MINIMA, TEMP_INICIAL * DECAIMENTO ** len(partidas))
            hero.temperatura = boy.temperatura = temp

            # alterna quem joga de brancas (le a ultima partida do Excel)
            hero_brancas = partidas[-1][2] != "Hero" if partidas else True
            ctx.brancas, ctx.pretas = (hero, boy) if hero_brancas else (boy, hero)
            brancas, pretas = ctx.brancas, ctx.pretas
            ctx.numero = len(partidas) + 1

            board, rb, motivo, lances_txt = jogar_partida(ctx, ui)

            if rb == 1:
                vencedor, texto = brancas.nome, "1-0"
            elif rb == -1:
                vencedor, texto = pretas.nome, "0-1"
            else:
                vencedor, texto = "Empate", "1/2-1/2"
            motivo_pt = MOTIVOS.get(motivo, motivo.lower())
            ctx.banner = ("Empate" if vencedor == "Empate" else f"{vencedor} venceu!", motivo_pt)
            ctx.msg = "Treinando Hero e Boy com esta partida..."
            ui.desenhar(ctx, force=True)
            ui.eventos()

            # APRENDIZADO: cada um guarda na sua memoria privada e treina a propria rede ----
            for agente, res in ((brancas, rb), (pretas, -rb)):
                idx = {1: 0, 0: 1, -1: 2}[res]
                agente.registrar_partida(res, ctx.numero)
                for antes, uci, *_ in agente.historico:
                    comp.registrar(antes, uci, idx)          # conhecimento publico
            erro_hero = hero.aprender()
            erro_boy = boy.aprender()
            # -------------------------------------------------------------------------------

            partidas.append((ctx.numero, datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                             brancas.nome, pretas.nome, texto, vencedor, motivo_pt, board.ply(),
                             erro_hero, erro_boy, lances_txt))
            ctx.recalcular()
            feitas += 1

            resumo = "empate" if vencedor == "Empate" else f"venceu: {vencedor}"
            print(f">>> Partida {ctx.numero}: {brancas.nome}(B) x {pretas.nome}(P)  {texto}  {resumo}  "
                  f"[{motivo_pt} | {board.ply()} meio-lances]  erro Hero: {erro_hero} | erro Boy: {erro_boy}")

            if ui.modo in ("lance", "partida") or feitas % 5 == 0:
                ctx.msg = "Salvando no Excel..."
                ui.desenhar(ctx, force=True)
                salvar_excel(hero, boy, comp, partidas)
            ctx.msg = ""
            ctx.numero = len(partidas) + 1
            ui.fim_partida(ctx)
    except (Sair, KeyboardInterrupt):
        print("\nEncerrando...")
    finally:
        ui.fechar()

    while not salvar_excel(hero, boy, comp, partidas):
        r = input("Feche o Excel e aperte Enter para tentar de novo (ou S para sair sem salvar): ")
        if r.strip().lower() == "s":
            break

    ctx.recalcular()
    print("\n=== PLACAR GERAL ===")
    print(f"Hero: {ctx.v_hero} | Boy: {ctx.v_boy} | Empates: {ctx.empates} | Total: {len(partidas)}")
    print(f"Memoria privada -> Hero: {len(hero.exp)} experiencias | Boy: {len(boy.exp)} experiencias")
    print(f"Salvo em: {ARQUIVO}")


if __name__ == "__main__":
    main()
